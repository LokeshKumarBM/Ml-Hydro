
from pathlib import Path
import csv
import warnings
from openpyxl import load_workbook

# This script is located in the ml folder.
ML_DIR = Path(__file__).resolve().parent

# Input and output folders.
RAW_DIR = ML_DIR / "data" / "raw" / "hydrogen_compatibility"
OUTPUT_DIR = ML_DIR / "data" / "processed"

WORKBOOKS = [
    "1100cia85tenfrafat.xlsx",
    "1211nib10frafat.xlsx",
    "3230san11frafat.xlsx",
]

# Sheet names containing the main fracture and fatigue test data.
TARGET_SHEETS = ("fracture", "fatigue")


def clean_value(value):
    """Convert Excel values into CSV-friendly values."""
    if value is None:
        return ""

    if isinstance(value, (int, float, bool)):
        return value

    return str(value).strip()


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    all_cells = []
    selected_cells = []

    for filename in WORKBOOKS:
        workbook_path = RAW_DIR / filename

        if not workbook_path.exists():
            print(f"WARNING: Workbook not found: {workbook_path}")
            continue

        print(f"\nReading: {filename}")

        # data_only=True reads cached formula results where available.
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Unknown extension is not supported"
            )
            workbook = load_workbook(
                workbook_path,
                read_only=True,
                data_only=True
            )

        for sheet in workbook.worksheets:
            sheet_name = sheet.title
            is_target_sheet = any(
                keyword in sheet_name.lower()
                for keyword in TARGET_SHEETS
            )

            nonempty_count = 0

            for row_number, row in enumerate(
                sheet.iter_rows(),
                start=1
            ):
                for column_number, cell in enumerate(
                    row,
                    start=1
                ):
                    value = clean_value(cell.value)

                    # Preserve non-empty cells and their original positions.
                    if value == "":
                        continue

                    record = {
                        "workbook": filename,
                        "sheet": sheet_name,
                        "row_number": row_number,
                        "column_number": column_number,
                        "column_letter": cell.column_letter,
                        "cell_value": value,
                    }

                    all_cells.append(record)
                    nonempty_count += 1

                    if is_target_sheet:
                        selected_cells.append(record)

            print(
                f"  {sheet_name}: "
                f"{nonempty_count} non-empty cells"
            )

        workbook.close()

    fields = [
        "workbook",
        "sheet",
        "row_number",
        "column_number",
        "column_letter",
        "cell_value",
    ]

    outputs = [
        (
            OUTPUT_DIR / "hydrogen_workbook_cells.csv",
            all_cells,
        ),
        (
            OUTPUT_DIR / "hydrogen_fracture_fatigue_cells.csv",
            selected_cells,
        ),
    ]

    for output_path, records in outputs:
        with output_path.open(
            "w",
            newline="",
            encoding="utf-8-sig"
        ) as csv_file:
            writer = csv.DictWriter(
                csv_file,
                fieldnames=fields
            )
            writer.writeheader()
            writer.writerows(records)

        print(f"\nSaved: {output_path}")
        print(f"Records: {len(records)}")

    if not all_cells:
        print(
            "\nNo data was extracted. Check the input folder "
            "and workbook filenames."
        )
    else:
        print("\nExtraction completed.")


if __name__ == "__main__":
    main()
