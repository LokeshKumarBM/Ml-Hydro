
from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "data" / "raw" / "hydrogen_compatibility"

WORKBOOKS = [
    "1100cia85tenfrafat.xlsx",
    "1211nib10frafat.xlsx",
    "3230san11frafat.xlsx",
]

for filename in WORKBOOKS:
    path = RAW_DIR / filename

    if not path.exists():
        print(f"\nWorkbook not found: {path}")
        continue

    workbook = load_workbook(
        path, read_only=True, data_only=True
    )

    for sheet in workbook.worksheets:
        if "fatigue" not in sheet.title.lower():
            continue

        print(f"\n{'=' * 70}")
        print(f"Workbook: {filename}")
        print(f"Sheet: {sheet.title}")
        print("=" * 70)

        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True), start=1
        ):
            values = [
                f"Column {i + 1}: {value}"
                for i, value in enumerate(row)
                if value is not None and str(value).strip()
            ]

            if values:
                print(f"Row {row_number}: " + " | ".join(values))

            if row_number >= 20:
                break

    workbook.close()
